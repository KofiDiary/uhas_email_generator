import os
import csv
import pandas as pd
from io import BytesIO
from email_logic import generate_email, resolve_duplicates


# ─── REQUIRED INPUT COLUMNS ───

# Maps expected logical field names to possible column header variations
COLUMN_ALIASES = {
    "student_id":           ["student no", "student number", "index number", "index no", "id"],
    "name":                 ["name", "full name", "fullname", "student name"],
    "programme":            ["program", "programme", "program offered", "programme offered"],
    "level":                ["level"],
    "level_of_admission":   ["level of admission", "admission level", "level admitted"],
}


# ─── COLUMN DETECTION ───

def detect_columns(df_columns):
    # Attempts to match actual DataFrame columns to expected logical field names
    # Returns a dict mapping logical name -> actual column name
    # Raises ValueError listing any required columns that could not be matched

    normalized = {col.strip().lower(): col for col in df_columns}
    mapping = {}
    missing = []

    for field, aliases in COLUMN_ALIASES.items():
        matched = None
        for alias in aliases:
            if alias.lower() in normalized:
                matched = normalized[alias.lower()]
                break
        if matched:
            mapping[field] = matched
        else:
            missing.append(field.replace("_", " ").title())

    if missing:
        raise ValueError(
            f"Could not detect the following required columns: {', '.join(missing)}.\n"
            f"Columns found in file: {', '.join(df_columns)}"
        )

    return mapping


# ─── FILE READING ───

def read_input_file(file_path_or_buffer, file_extension=None):
    # Reads a CSV or Excel file into a DataFrame
    # file_path_or_buffer can be a file path string or a file-like object
    # file_extension should be ".csv" or ".xlsx" / ".xls" when using a buffer
    # Returns a cleaned DataFrame

    if isinstance(file_path_or_buffer, str):
        ext = os.path.splitext(file_path_or_buffer)[1].lower()
    else:
        ext = (file_extension or "").lower()

    if ext == ".csv":
        df = pd.read_csv(file_path_or_buffer, dtype=str)
    elif ext in (".xlsx", ".xls"):
        df = pd.read_excel(file_path_or_buffer, dtype=str)
    else:
        raise ValueError(
            f"Unsupported file type '{ext}'. Please upload a .csv or .xlsx file."
        )

    # Strip whitespace from all column headers
    df.columns = [str(c).strip() for c in df.columns]

    # Drop completely empty rows
    df.dropna(how="all", inplace=True)

    # Reset index after dropping rows
    df.reset_index(drop=True, inplace=True)

    return df


# ─── ROW PROCESSING ───

def process_dataframe(df, col_map, student_type):
    # Iterates over every row in the DataFrame and generates email data
    # col_map: dict returned by detect_columns()
    # student_type: "Regular" or "Sandwich"
    # Returns a list of result dicts, one per row

    results = []

    for idx, row in df.iterrows():
        row_number = idx + 2  # Account for header row in Excel/CSV

        student_id        = str(row[col_map["student_id"]]).strip()
        full_name         = str(row[col_map["name"]]).strip()
        programme         = str(row[col_map["programme"]]).strip()
        level_of_admission = str(row[col_map["level_of_admission"]]).strip()

        # Skip rows where critical fields are empty or NaN
        if any(v.lower() in ("", "nan", "none") for v in [student_id, full_name, programme, level_of_admission]):
            results.append({
                "row_number": row_number,
                "student_id": student_id,
                "email": "",
                "first_name": "",
                "last_name": "",
                "display_name": "",
                "department": "",
                "error": "One or more required fields are empty. Row skipped.",
                "skipped": True,
            })
            continue

        # Generate email for this row
        email_data = generate_email(
            full_name=full_name,
            student_id=student_id,
            programme=programme,
            level_of_admission=level_of_admission,
            student_type=student_type,
        )

        results.append({
            "row_number": row_number,
            "student_id": student_id,
            "email": email_data["email"],
            "first_name": email_data["first_name"],
            "last_name": email_data["last_name"],
            "display_name": email_data["display_name"],
            "department": email_data["department"],
            "error": email_data["error"],
            "skipped": False,
        })

    # Resolve duplicate emails across all valid rows
    all_emails = [r["email"] for r in results]
    resolved_emails = resolve_duplicates(all_emails)
    for i, result in enumerate(results):
        result["email"] = resolved_emails[i]
        # Keep alternate email in sync with resolved email
        result["alternate_email"] = resolved_emails[i]

    return results


# ─── OUTPUT FILE BUILDING ───

def build_output_dataframe(results):
    # Builds the final output DataFrame from processed results
    # Follows the exact 16-column format required for import
    # Only populates columns that should have data, rest are left blank

    rows = []

    for r in results:
        rows.append({
            "Username":                 r["email"],
            "First name":               r["first_name"],
            "Last name":                r["last_name"],
            "Display name":             r["display_name"],
            "Job title":                "Student" if not r["skipped"] else "",
            "DEPARTMENT":               r["department"],
            "Office number":            "",
            "Office phone":             "",
            "Mobile phone":             "",
            "Fax Number":               r["student_id"] if not r["skipped"] else "",
            "Alternate email address":  r.get("alternate_email", ""),
            "Address":                  "",
            "City":                     "Ho" if not r["skipped"] else "",
            "State or province":        "Volta" if not r["skipped"] else "",
            "ZIP or postal code":       "",
            "Country or region":        "Ghana" if not r["skipped"] else "",
        })

    return pd.DataFrame(rows)


# ─── EXPORT TO EXCEL ───

def export_to_excel(output_df):
    # Writes the output DataFrame to a formatted Excel file in memory
    # Returns a BytesIO buffer ready for download

    buffer = BytesIO()

    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        output_df.to_excel(writer, index=False, sheet_name="Emails")

        workbook  = writer.book
        worksheet = writer.sheets["Emails"]

        # Import openpyxl styles here to keep the top-level imports clean
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        # Header style
        header_font    = Font(bold=True, color="FFFFFF", size=11)
        header_fill    = PatternFill("solid", fgColor="1F4E79")
        header_align   = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border    = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        # Apply header styles
        for col_idx, cell in enumerate(worksheet[1], start=1):
            cell.font      = header_font
            cell.fill      = header_fill
            cell.alignment = header_align
            cell.border    = thin_border

        # Apply row styles and auto-fit columns
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="center")
                cell.border    = thin_border
                # Light alternating row fill
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="D6E4F0")

        # Auto-fit column widths based on content
        for col_idx, col_cells in enumerate(worksheet.columns, start=1):
            max_length = 0
            col_letter = get_column_letter(col_idx)
            for cell in col_cells:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except Exception:
                    pass
            # Add padding, cap at 60 characters wide
            adjusted_width = min(max_length + 4, 60)
            worksheet.column_dimensions[col_letter].width = adjusted_width

        # Freeze the header row
        worksheet.freeze_panes = "A2"

    buffer.seek(0)
    return buffer


# ─── EXPORT TO CSV ───

def export_to_csv(output_df):
    # Writes the output DataFrame to a CSV file in memory
    # Returns a BytesIO buffer ready for download

    buffer = BytesIO()
    output_df.to_csv(buffer, index=False, encoding="utf-8-sig")
    buffer.seek(0)
    return buffer


# ─── ERROR REPORT BUILDING ───

def build_error_report(results):
    # Builds a summary of rows that had errors or were skipped
    # Returns a list of dicts for display in the UI

    error_rows = [
        {
            "Row":          r["row_number"],
            "Student ID":   r["student_id"],
            "Issue":        r["error"] or "Unknown error",
        }
        for r in results
        if r.get("error") or r.get("skipped")
    ]

    return error_rows


# ─── MAIN PIPELINE ───

def run_pipeline(file_path_or_buffer, file_extension, student_type):
    # Full processing pipeline from raw file to output buffers
    # Returns a dict with keys:
    #   output_df        - the final output DataFrame
    #   excel_buffer     - BytesIO for Excel download
    #   csv_buffer       - BytesIO for CSV download
    #   error_report     - list of dicts describing problem rows
    #   total_rows       - int total rows processed
    #   success_count    - int rows with emails generated
    #   skipped_count    - int rows skipped due to errors

    # Read input file
    df = read_input_file(file_path_or_buffer, file_extension)

    # Detect required columns
    col_map = detect_columns(df.columns.tolist())

    # Process all rows
    results = process_dataframe(df, col_map, student_type)

    # Build output DataFrame
    output_df = build_output_dataframe(results)

    # Build download buffers
    excel_buffer = export_to_excel(output_df)
    csv_buffer   = export_to_csv(output_df)

    # Build error report
    error_report  = build_error_report(results)
    success_count = sum(1 for r in results if not r.get("skipped") and not r.get("error"))
    skipped_count = len(results) - success_count

    return {
        "output_df":     output_df,
        "excel_buffer":  excel_buffer,
        "csv_buffer":    csv_buffer,
        "error_report":  error_report,
        "total_rows":    len(results),
        "success_count": success_count,
        "skipped_count": skipped_count,
    }